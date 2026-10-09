"""Generate Revue-pages-Riyad-Zaer-Gardens.xlsx for Boostetudes review."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pathlib import Path

BASE = "https://riyadzaergardens.com"

HEADERS = [
    "Page",
    "Link",
    "Status développement",
    "Avis Boostetudes",
    "Remarques",
    "Avis Boostetudes in Mobile",
    "Remarques mobile",
]

ROWS = [
    (
        "Accueil (FR)",
        f"{BASE}/",
        "Prêt — showcase",
        "",
        "Hero photo, sections projet / logements / commerces / lieu / contact. Formulaire home branché vers /merci.",
        "",
        "",
    ),
    (
        "Le projet (FR)",
        f"{BASE}/le-projet",
        "Prêt — showcase",
        "",
        "Hero full-bleed + facts, conception, cadre de vie, prestations, syndic, fiche, CTA photo.",
        "",
        "",
    ),
    (
        "Appartements (FR)",
        f"{BASE}/appartements",
        "Prêt — showcase",
        "",
        "Hero photo, typologies, inclus, finitions, prix / aide, CTA photo.",
        "",
        "",
    ),
    (
        "Fonds de commerce (FR)",
        f"{BASE}/fonds-de-commerce",
        "Prêt — showcase",
        "",
        "Hero photo, offre, raisons, investissement, CTA photo.",
        "",
        "",
    ),
    (
        "Localisation (FR)",
        f"{BASE}/localisation",
        "Prêt — showcase",
        "",
        "Hero photo, accès, carte, quartier, CTA visite photo.",
        "",
        "",
    ),
    (
        "Contact (FR)",
        f"{BASE}/contact",
        "Prêt — showcase",
        "",
        "Hero brand + facts, scène visite, formulaire + aside, more links.",
        "",
        "",
    ),
    (
        "FAQ (FR)",
        f"{BASE}/faq",
        "Prêt — showcase",
        "",
        "Hero + facts, accordéons par thème, CTA photo, more links.",
        "",
        "",
    ),
    (
        "Blog — listing (FR)",
        f"{BASE}/blog",
        "Prêt — showcase",
        "",
        "Hero + article à la une, liste éditoriale photos, CTA, more links.",
        "",
        "",
    ),
    (
        "Blog — articles (FR)",
        f"{BASE}/blog/[slug]",
        "Prêt",
        "",
        "5 articles. Hero photo par article, corps, articles liés, CTA. Détail → feuille « Articles blog ».",
        "",
        "",
    ),
    (
        "La Manoussa / Promoteur (FR)",
        f"{BASE}/la-manoussa",
        "Prêt — showcase",
        "",
        "Page promoteur refaite : hero brand, rôle, projet, équipes, CTA.",
        "",
        "",
    ),
    (
        "Mentions légales (FR)",
        f"{BASE}/mentions-legales",
        "Prêt",
        "",
        "Hero, éditeur, PI, confidentialité (section sombre). Ancre #confidentialite.",
        "",
        "",
    ),
    (
        "Merci (FR)",
        f"{BASE}/merci",
        "Prêt",
        "",
        "Confirmation post-formulaire (noindex). Hero + liens continuer.",
        "",
        "",
    ),
    (
        "404 (FR)",
        f"{BASE}/404",
        "Prêt",
        "",
        "Hero photo + CTAs + liens continuer.",
        "",
        "",
    ),
    (
        "Guides (FR)",
        f"{BASE}/guides",
        "Redirection 301 → /blog",
        "",
        "Ancienne URL redirigée vers le blog.",
        "",
        "",
    ),
    (
        "Accueil (AR)",
        f"{BASE}/ar/",
        "Prêt — showcase",
        "",
        "Miroir RTL de l’accueil FR.",
        "",
        "",
    ),
    (
        "المشروع / Le projet (AR)",
        f"{BASE}/ar/le-projet",
        "Prêt — showcase",
        "",
        "Miroir RTL showcase projet.",
        "",
        "",
    ),
    (
        "الشقق / Appartements (AR)",
        f"{BASE}/ar/appartements",
        "Prêt — showcase",
        "",
        "Miroir RTL showcase appartements.",
        "",
        "",
    ),
    (
        "المحلات / Fonds de commerce (AR)",
        f"{BASE}/ar/fonds-de-commerce",
        "Prêt — showcase",
        "",
        "Miroir RTL showcase commerces.",
        "",
        "",
    ),
    (
        "الموقع / Localisation (AR)",
        f"{BASE}/ar/localisation",
        "Prêt — showcase",
        "",
        "Miroir RTL showcase localisation.",
        "",
        "",
    ),
    (
        "اتصل بنا / Contact (AR)",
        f"{BASE}/ar/contact",
        "Prêt — showcase",
        "",
        "Miroir RTL contact + formulaire.",
        "",
        "",
    ),
    (
        "الأسئلة الشائعة / FAQ (AR)",
        f"{BASE}/ar/faq",
        "Prêt — showcase",
        "",
        "Miroir RTL FAQ.",
        "",
        "",
    ),
    (
        "المدونة / Blog listing (AR)",
        f"{BASE}/ar/blog",
        "Prêt — showcase",
        "",
        "Miroir RTL blog listing.",
        "",
        "",
    ),
    (
        "Blog — articles (AR)",
        f"{BASE}/ar/blog/[slug]",
        "Prêt",
        "",
        "Articles AR miroir des FR. Détail → feuille « Articles blog ».",
        "",
        "",
    ),
    (
        "المطوّر / La Manoussa (AR)",
        f"{BASE}/ar/la-manoussa",
        "Prêt — showcase",
        "",
        "Miroir RTL promoteur.",
        "",
        "",
    ),
    (
        "إشارات قانونية / Mentions (AR)",
        f"{BASE}/ar/mentions-legales",
        "Prêt",
        "",
        "Miroir RTL mentions + confidentialité.",
        "",
        "",
    ),
    (
        "شكراً / Merci (AR)",
        f"{BASE}/ar/merci",
        "Prêt",
        "",
        "Confirmation AR (noindex).",
        "",
        "",
    ),
    (
        "404 (AR)",
        f"{BASE}/ar/404",
        "Prêt",
        "",
        "404 AR.",
        "",
        "",
    ),
]

ARTICLES = [
    ("Aide au logement (FR)", f"{BASE}/blog/aide-au-logement-maroc", "Prêt", "", "Article guide.", "", ""),
    ("F3 ou F4 (FR)", f"{BASE}/blog/f3-ou-f4-comment-choisir", "Prêt", "", "Article guide.", "", ""),
    ("Acheter à Aïn Aouda (FR)", f"{BASE}/blog/acheter-a-ain-aouda", "Prêt", "", "Article guide.", "", ""),
    ("Investir neuf près de Rabat (FR)", f"{BASE}/blog/investir-neuf-pres-de-rabat", "Prêt", "", "Article guide.", "", ""),
    ("Visite du site (FR)", f"{BASE}/blog/visite-du-site-pourquoi-venir", "Prêt", "", "Article guide.", "", ""),
    ("Aide au logement (AR)", f"{BASE}/ar/blog/aide-au-logement-maroc", "Prêt", "", "Miroir AR.", "", ""),
    ("F3 ou F4 (AR)", f"{BASE}/ar/blog/f3-ou-f4-comment-choisir", "Prêt", "", "Miroir AR.", "", ""),
    ("Acheter à Aïn Aouda (AR)", f"{BASE}/ar/blog/acheter-a-ain-aouda", "Prêt", "", "Miroir AR.", "", ""),
    ("Investir neuf (AR)", f"{BASE}/ar/blog/investir-neuf-pres-de-rabat", "Prêt", "", "Miroir AR.", "", ""),
    ("Visite du site (AR)", f"{BASE}/ar/blog/visite-du-site-pourquoi-venir", "Prêt", "", "Miroir AR.", "", ""),
]

WIDTHS = [34, 52, 24, 22, 58, 26, 40]


def style_sheet(ws, rows):
    header_fill = PatternFill("solid", fgColor="002D2D")
    header_font = Font(name="Calibri", bold=True, color="F4EEE2", size=11)
    cell_font = Font(name="Calibri", size=10, color="002D2D")
    link_font = Font(name="Calibri", size=10, color="0A66C2", underline="single")
    thin = Border(
        left=Side(style="thin", color="C8B568"),
        right=Side(style="thin", color="C8B568"),
        top=Side(style="thin", color="C8B568"),
        bottom=Side(style="thin", color="C8B568"),
    )
    alt_fill = PatternFill("solid", fgColor="FCFAF6")
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(wrap_text=True, vertical="center", horizontal="center")
    status_fills = {
        "showcase": PatternFill("solid", fgColor="D4EDDA"),
        "ready": PatternFill("solid", fgColor="E8F0FE"),
        "redir": PatternFill("solid", fgColor="FFF3CD"),
    }

    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin

    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{len(rows) + 1}"

    for r_idx, row in enumerate(rows, 2):
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, val)
            cell.font = cell_font
            cell.border = thin
            cell.alignment = wrap
            if r_idx % 2 == 0:
                cell.fill = alt_fill
            if c_idx == 2 and val:
                cell.font = link_font
                href = val.replace("[slug]", "aide-au-logement-maroc")
                cell.hyperlink = href
            if c_idx == 3:
                cell.alignment = center
                if "showcase" in val:
                    cell.fill = status_fills["showcase"]
                elif "Redirection" in val:
                    cell.fill = status_fills["redir"]
                elif val == "Prêt":
                    cell.fill = status_fills["ready"]
        ws.row_dimensions[r_idx].height = 48

    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    avis = DataValidation(
        type="list",
        formula1='"OK,À revoir,Bloquant,Non testé"',
        allow_blank=True,
    )
    avis.prompt = "Avis Boostetudes"
    ws.add_data_validation(avis)
    avis.add(f"D2:D{len(rows) + 1}")
    avis.add(f"F2:F{len(rows) + 1}")


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Revue pages"
    style_sheet(ws, ROWS)

    ws2 = wb.create_sheet("Légende & consignes")
    ws2["A1"] = "Revue pages — Riyad Zaer Gardens (site-v3)"
    ws2["A1"].font = Font(name="Calibri", bold=True, size=14, color="002D2D")
    ws2.merge_cells("A1:B1")

    cream = PatternFill("solid", fgColor="F4EEE2")
    legend = [
        ("", ""),
        ("Colonnes", "Mode d’emploi"),
        ("Page", "Nom de la page (FR / AR)."),
        ("Link", "URL de production (cliquable). Adapter le domaine si préprod."),
        ("Status développement", "État côté build actuel (déjà rempli)."),
        ("Avis Boostetudes", "À remplir : OK / À revoir / Bloquant / Non testé (liste déroulante)."),
        ("Remarques", "Notes dev déjà renseignées — Boostetudes peut compléter."),
        ("Avis Boostetudes in Mobile", "Même grille, revue mobile (phone / tablette)."),
        ("Remarques mobile", "Bugs ou retours spécifiques mobile."),
        ("", ""),
        ("Status — valeurs", "Signification"),
        ("Prêt — showcase", "Page complète au standard actuel (hero photo, sections, CTA)."),
        ("Prêt", "Page fonctionnelle et soignée, format plus court."),
        ("Redirection 301 → /blog", "URL conservée, redirige vers le blog."),
        ("", ""),
        ("Périmètre", "FR + AR. 5 articles blog FR + 5 AR (feuille dédiée)."),
        ("Date export", "2026-10-07"),
        ("Domaine utilisé", BASE),
        ("Dev local", "http://localhost:4330"),
    ]
    for i, (a, b) in enumerate(legend, 3):
        ca = ws2.cell(i, 1, a)
        cb = ws2.cell(i, 2, b)
        ca.font = Font(name="Calibri", size=10, color="002D2D")
        cb.font = Font(name="Calibri", size=10, color="002D2D")
        if a in ("Colonnes", "Status — valeurs"):
            ca.font = Font(name="Calibri", bold=True, size=12, color="002D2D")
            ca.fill = cream
            cb.fill = cream
        elif a:
            ca.font = Font(name="Calibri", bold=True, size=10, color="002D2D")
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 88

    ws3 = wb.create_sheet("Articles blog")
    style_sheet(ws3, ARTICLES)

    out = Path(__file__).resolve().parents[1] / "Revue-pages-Riyad-Zaer-Gardens.xlsx"
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
